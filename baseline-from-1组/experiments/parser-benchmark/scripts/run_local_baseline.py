"""Run the installed local parser baseline against project-authored fixtures.

This is an experiment runner, not a product parser. It intentionally records
missing optional adapters (for example OCR) instead of installing or guessing.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import re
import sys
import time
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples" / "self-authored"
RESULTS = ROOT / "results" / "local-baseline"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def version(package: str) -> str | None:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return None


def annotation(sample_id: str) -> dict[str, Any]:
    path = SAMPLES / "annotations" / f"{sample_id}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def query_hits(text: str, queries: list[str]) -> dict[str, bool]:
    return {query: query.casefold() in text.casefold() for query in queries}


def parse_pdf(sample_id: str, path: Path) -> dict[str, Any]:
    import fitz  # PyMuPDF

    started = time.perf_counter()
    document = fitz.open(path)
    pages: list[dict[str, Any]] = []
    all_text: list[str] = []
    total_blocks = 0
    total_images = 0
    for index, page in enumerate(document, start=1):
        text = page.get_text("text")
        blocks = page.get_text("blocks", sort=False)
        images = page.get_images(full=True)
        total_blocks += len(blocks)
        total_images += len(images)
        all_text.append(text)
        pages.append(
            {
                "physical_page": index,
                "text_chars": len(text),
                "block_count": len(blocks),
                "image_count": len(images),
                "bbox_sample": list(blocks[0][:4]) if blocks else None,
            }
        )
    text = "\n".join(all_text)
    result = {
        "parser": "PyMuPDF",
        "parser_version": version("PyMuPDF"),
        "status": "succeeded",
        "sample_id": sample_id,
        "format": "PDF",
        "pages": pages,
        "page_count": len(document),
        "text_chars": len(text),
        "block_count": total_blocks,
        "image_count": total_images,
        "text_layer_present": bool(text.strip()),
        "ocr_status": "not_run",
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "text_preview": text[:1000],
    }
    result["query_hits"] = query_hits(text, annotation(sample_id).get("fixed_queries", []))
    document.close()
    return result


def parse_docx(sample_id: str, path: Path) -> dict[str, Any]:
    from docx import Document

    started = time.perf_counter()
    document = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text
        paragraphs.append(
            {
                "index": index,
                "style": paragraph.style.name if paragraph.style else None,
                "text": text,
            }
        )
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        tables.append(
            {
                "index": table_index,
                "rows": len(table.rows),
                "columns": max((len(row.cells) for row in table.rows), default=0),
                "text": [[cell.text for cell in row.cells] for row in table.rows],
            }
        )
    text = "\n".join(item["text"] for item in paragraphs)
    result = {
        "parser": "python-docx",
        "parser_version": version("python-docx"),
        "status": "succeeded",
        "sample_id": sample_id,
        "format": "DOCX",
        "paragraph_count": len(paragraphs),
        "heading_count": sum(
            1 for item in paragraphs if str(item["style"]).lower().startswith("heading")
        ),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "visual_page_numbers": "not_available_from_python_docx",
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "text_preview": text[:1000],
    }
    result["query_hits"] = query_hits(text, annotation(sample_id).get("fixed_queries", []))
    return result


def parse_markdown(sample_id: str, path: Path) -> dict[str, Any]:
    from markdown_it import MarkdownIt

    started = time.perf_counter()
    source = path.read_text(encoding="utf-8")
    tokens = MarkdownIt("commonmark").parse(source)
    headings = []
    for index, token in enumerate(tokens):
        if token.type != "heading_open":
            continue
        inline = tokens[index + 1] if index + 1 < len(tokens) else None
        headings.append(
            {
                "level": int(token.tag[1:]),
                "text": inline.content if inline and inline.type == "inline" else "",
                "map": token.map,
            }
        )
    result = {
        "parser": "markdown-it-py",
        "parser_version": version("markdown-it-py"),
        "status": "succeeded",
        "sample_id": sample_id,
        "format": "Markdown",
        "line_count": len(source.splitlines()),
        "token_count": len(tokens),
        "heading_count": len(headings),
        "headings": headings,
        "commonmark_mode": True,
        "gfm_extensions": "not_enabled_in_baseline",
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "text_preview": source[:1000],
    }
    result["query_hits"] = query_hits(source, annotation(sample_id).get("fixed_queries", []))
    return result


def parse_txt(sample_id: str, path: Path) -> dict[str, Any]:
    started = time.perf_counter()
    raw = path.read_bytes()
    encoding = "utf-8"
    decode_error = None
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError as error:
        decode_error = str(error)
        encoding = "gb18030"
        text = raw.decode(encoding)
    lines = text.splitlines()
    paragraphs = [part for part in re.split(r"\n\s*\n", text) if part.strip()]
    result = {
        "parser": "explicit-encoding-baseline",
        "parser_version": "stdlib",
        "status": "succeeded",
        "sample_id": sample_id,
        "format": "TXT",
        "encoding": encoding,
        "initial_utf8_decode_error": decode_error,
        "line_count": len(lines),
        "paragraph_count_by_blank_line": len(paragraphs),
        "crlf_count": raw.count(b"\r\n"),
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "text_preview": text[:1000],
    }
    result["query_hits"] = query_hits(text, annotation(sample_id).get("fixed_queries", []))
    return result


def samples_for(kind: str) -> list[tuple[str, Path]]:
    groups = {
        "pdf": [("SELF-001", SAMPLES / "pdf" / "SELF-001-single-column.pdf"),
                ("SELF-002", SAMPLES / "pdf" / "SELF-002-two-column.pdf"),
                ("SELF-003", SAMPLES / "pdf" / "SELF-003-table-formula.pdf"),
                ("SELF-004", SAMPLES / "pdf" / "SELF-004-scanned.pdf")],
        "docx": [("SELF-005", SAMPLES / "docx" / "SELF-005-structured.docx")],
        "markdown": [("SELF-008", SAMPLES / "markdown" / "SELF-008-complex.md")],
        "txt": [("SELF-010", SAMPLES / "txt" / "SELF-010-utf8.txt"),
                ("SELF-011", SAMPLES / "txt" / "SELF-011-gb18030.txt")],
    }
    if kind == "all":
        return [item for group in groups.values() for item in group]
    return groups[kind]


def run(kind: str) -> dict[str, Any]:
    RESULTS.mkdir(parents=True, exist_ok=True)
    outputs = []
    for sample_id, path in samples_for(kind):
        if not path.exists():
            outputs.append({"sample_id": sample_id, "status": "missing", "path": str(path)})
            continue
        try:
            if path.suffix.lower() == ".pdf":
                result = parse_pdf(sample_id, path)
            elif path.suffix.lower() == ".docx":
                result = parse_docx(sample_id, path)
            elif path.suffix.lower() == ".md":
                result = parse_markdown(sample_id, path)
            else:
                result = parse_txt(sample_id, path)
            result["source_path"] = str(path.relative_to(ROOT))
            result["source_sha256"] = sha256(path)
        except Exception as error:  # keep one failed sample from hiding others
            result = {
                "sample_id": sample_id,
                "status": "failed",
                "source_path": str(path.relative_to(ROOT)),
                "error_type": type(error).__name__,
                "error": str(error),
            }
        output_path = RESULTS / f"{sample_id}.json"
        output_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        outputs.append(result)
    summary = {
        "experiment_id": "TECH-DOC-001-R01-local-baseline",
        "date": time.strftime("%Y-%m-%d"),
        "python": sys.version,
        "kind": kind,
        "outputs": outputs,
        "tool_availability": {
            "PyMuPDF": version("PyMuPDF"),
            "pdfplumber": version("pdfplumber"),
            "docling": version("docling"),
            "unstructured": version("unstructured"),
            "python-docx": version("python-docx"),
            "markdown-it-py": version("markdown-it-py"),
        },
        "boundary": "Local baseline only; no OCR, cloud service, vector retrieval, or QA evaluation.",
    }
    (RESULTS / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--kind", choices=["all", "pdf", "docx", "markdown", "txt"], default="all")
    args = parser.parse_args()
    summary = run(args.kind)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
