"""Run non-OCR document structure checks on local authorized fixtures.

This experiment deliberately excludes OCR. Scanned PDFs are recorded as an
out-of-scope/no-text-layer case and are not scored as parser failures or wins.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import re
import sys
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples" / "self-authored"
RESULTS = ROOT / "results" / "text-structure"

SAMPLE_FILES = {
    "SELF-001": SAMPLES / "pdf" / "SELF-001-single-column.pdf",
    "SELF-002": SAMPLES / "pdf" / "SELF-002-two-column.pdf",
    "SELF-003": SAMPLES / "pdf" / "SELF-003-table-formula.pdf",
    "SELF-004": SAMPLES / "pdf" / "SELF-004-scanned.pdf",
    "SELF-005": SAMPLES / "docx" / "SELF-005-structured.docx",
    "SELF-006": SAMPLES / "docx" / "SELF-006-meeting-notes.docx",
    "SELF-007": SAMPLES / "docx" / "SELF-007-complex.docx",
    "SELF-008": SAMPLES / "markdown" / "SELF-008-complex.md",
    "SELF-009": SAMPLES / "markdown" / "SELF-009-heading-levels.md",
    "SELF-010": SAMPLES / "txt" / "SELF-010-utf8.txt",
    "SELF-011": SAMPLES / "txt" / "SELF-011-gb18030.txt",
}

ANNOTATIONS = SAMPLES / "annotations"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def version(package: str) -> str | None:
    try:
        return importlib.metadata.version(package)
    except importlib.metadata.PackageNotFoundError:
        return None


def annotation(sample_id: str) -> dict[str, Any]:
    path = ANNOTATIONS / f"{sample_id}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def query_hits(text: str, queries: list[str]) -> dict[str, bool]:
    folded = text.casefold()
    return {query: query.casefold() in folded for query in queries}


def text_pdf(sample_id: str, path: Path) -> dict[str, Any]:
    import fitz
    import pdfplumber

    started = time.perf_counter()
    pymupdf_pages = []
    text_parts: list[str] = []
    with fitz.open(path) as document:
        for number, page in enumerate(document, start=1):
            text = page.get_text("text", sort=False)
            blocks = page.get_text("blocks", sort=False)
            words = page.get_text("words", sort=False)
            text_parts.append(text)
            pymupdf_pages.append(
                {
                    "physical_page": number,
                    "text_chars": len(text),
                    "block_count": len(blocks),
                    "word_count": len(words),
                    "bbox_sample": list(words[0][:4]) if words else None,
                }
            )
    full_text = "\n".join(text_parts)

    pdfplumber_pages = []
    with pdfplumber.open(path) as document:
        for number, page in enumerate(document.pages, start=1):
            extracted = page.extract_text() or ""
            words = page.extract_words() or []
            pdfplumber_pages.append(
                {
                    "physical_page": number,
                    "text_chars": len(extracted),
                    "word_count": len(words),
                    "table_candidate_count": len(page.find_tables()),
                }
            )

    result = {
        "experiment_id": "TECH-DOC-007-R01-text-structure",
        "sample_id": sample_id,
        "format": "PDF",
        "status": "scored" if full_text.strip() else "out_of_scope_no_text_layer",
        "ocr": {"enabled": False, "status": "not_run"},
        "pymupdf": {"version": version("PyMuPDF"), "pages": pymupdf_pages},
        "pdfplumber": {"version": version("pdfplumber"), "pages": pdfplumber_pages},
        "page_count": len(pymupdf_pages),
        "text_chars": len(full_text),
        "query_hits": query_hits(full_text, annotation(sample_id).get("fixed_queries", [])),
        "text_preview": full_text[:1500],
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "boundary": "Text-layer/structure experiment only; OCR quality is out of scope.",
    }
    return result


def docx(sample_id: str, path: Path) -> dict[str, Any]:
    from docx import Document

    started = time.perf_counter()
    document = Document(path)
    paragraphs = [
        {"index": i, "style": p.style.name if p.style else None, "text": p.text}
        for i, p in enumerate(document.paragraphs, start=1)
    ]
    tables = [
        {
            "index": i,
            "rows": len(table.rows),
            "columns": max((len(row.cells) for row in table.rows), default=0),
            "text": [[cell.text for cell in row.cells] for row in table.rows],
        }
        for i, table in enumerate(document.tables, start=1)
    ]
    text = "\n".join(item["text"] for item in paragraphs)
    return {
        "experiment_id": "TECH-DOC-007-R01-text-structure",
        "sample_id": sample_id,
        "format": "DOCX",
        "status": "scored",
        "ocr": {"enabled": False, "status": "not_applicable"},
        "parser": "python-docx",
        "parser_version": version("python-docx"),
        "paragraph_count": len(paragraphs),
        "heading_count": sum(1 for p in paragraphs if str(p["style"]).lower().startswith("heading")),
        "table_count": len(tables),
        "paragraphs": paragraphs,
        "tables": tables,
        "visual_page_numbers": "not_available_from_python_docx",
        "query_hits": query_hits(text, annotation(sample_id).get("fixed_queries", [])),
        "text_preview": text[:1500],
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "boundary": "DOCX logical structure only; final visual pagination is out of scope.",
    }


def markdown(sample_id: str, path: Path) -> dict[str, Any]:
    from markdown_it import MarkdownIt

    started = time.perf_counter()
    source = path.read_text(encoding="utf-8")
    tokens = MarkdownIt("commonmark").parse(source)
    headings = []
    for index, token in enumerate(tokens):
        if token.type != "heading_open":
            continue
        inline = tokens[index + 1] if index + 1 < len(tokens) else None
        headings.append(
            {
                "level": int(token.tag[1:]),
                "text": inline.content if inline and inline.type == "inline" else "",
                "source_line_map": token.map,
            }
        )
    return {
        "experiment_id": "TECH-DOC-007-R01-text-structure",
        "sample_id": sample_id,
        "format": "Markdown",
        "status": "scored",
        "ocr": {"enabled": False, "status": "not_applicable"},
        "parser": "markdown-it-py",
        "parser_version": version("markdown-it-py"),
        "commonmark_mode": True,
        "gfm_extensions": "not_enabled",
        "line_count": len(source.splitlines()),
        "token_count": len(tokens),
        "heading_count": len(headings),
        "headings": headings,
        "query_hits": query_hits(source, annotation(sample_id).get("fixed_queries", [])),
        "text_preview": source[:1500],
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "boundary": "CommonMark structure only; GFM/footnote/math extensions are out of scope.",
    }


def txt(sample_id: str, path: Path) -> dict[str, Any]:
    started = time.perf_counter()
    raw = path.read_bytes()
    encoding = "utf-8"
    decode_error = None
    try:
        text = raw.decode(encoding)
    except UnicodeDecodeError as error:
        decode_error = str(error)
        encoding = "gb18030"
        text = raw.decode(encoding)
    paragraphs = [part for part in re.split(r"\n\s*\n", text) if part.strip()]
    return {
        "experiment_id": "TECH-DOC-007-R01-text-structure",
        "sample_id": sample_id,
        "format": "TXT",
        "status": "scored",
        "ocr": {"enabled": False, "status": "not_applicable"},
        "parser": "explicit-encoding-baseline",
        "parser_version": "stdlib",
        "encoding": encoding,
        "initial_utf8_decode_error": decode_error,
        "line_count": len(text.splitlines()),
        "paragraph_count_by_blank_line": len(paragraphs),
        "crlf_count": raw.count(b"\r\n"),
        "query_hits": query_hits(text, annotation(sample_id).get("fixed_queries", [])),
        "text_preview": text[:1500],
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
        "boundary": "Encoding and line/blank-line structure only; no semantic paragraph inference.",
    }


def run(sample_ids: list[str]) -> dict[str, Any]:
    RESULTS.mkdir(parents=True, exist_ok=True)
    outputs = []
    for sample_id in sample_ids:
        path = SAMPLE_FILES[sample_id]
        try:
            if path.suffix.lower() == ".pdf":
                result = text_pdf(sample_id, path)
            elif path.suffix.lower() == ".docx":
                result = docx(sample_id, path)
            elif path.suffix.lower() == ".md":
                result = markdown(sample_id, path)
            else:
                result = txt(sample_id, path)
            result["source_path"] = str(path.relative_to(ROOT))
            result["source_sha256"] = sha256(path)
        except Exception as error:
            result = {
                "experiment_id": "TECH-DOC-007-R01-text-structure",
                "sample_id": sample_id,
                "status": "failed",
                "error_type": type(error).__name__,
                "error": str(error),
                "source_path": str(path.relative_to(ROOT)),
                "source_sha256": sha256(path),
                "ocr": {"enabled": False, "status": "not_run"},
            }
        (RESULTS / f"{sample_id}.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        outputs.append(result)
    summary = {
        "experiment_id": "TECH-DOC-007-R01-text-structure",
        "date": time.strftime("%Y-%m-%d"),
        "python": sys.version,
        "ocr_enabled": False,
        "samples": sample_ids,
        "outputs": outputs,
        "boundary": "No OCR, cloud service, vector retrieval, or QA evaluation.",
    }
    (RESULTS / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sample", action="append", choices=sorted(SAMPLE_FILES), dest="samples")
    args = parser.parse_args()
    print(json.dumps(run(args.samples or list(SAMPLE_FILES)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
